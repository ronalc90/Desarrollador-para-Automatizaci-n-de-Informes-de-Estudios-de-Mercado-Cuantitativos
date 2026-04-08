"""Actualización de gráficos y Excel embebido dentro de un .pptx.

Este módulo es el corazón técnico del motor. Cada gráfico en un PPT
tiene dos fuentes de verdad que deben mantenerse sincronizadas:

1. El XML del gráfico en ``ppt/charts/chartN.xml`` (lo que PowerPoint
   dibuja en pantalla).
2. El Excel embebido en ``ppt/embeddings/Microsoft_Excel_Sheet*.xlsx``
   (lo que se abre cuando el usuario hace click en *Editar datos*).

Para actualizar ambas cosas de forma consistente:

- Se construye un ``CategoryChartData`` a partir del DataFrame y se
  invoca ``chart.replace_data(...)``. Esto resuelve casi todos los
  casos simples y preserva formato visual (colores, fuentes, estilos).
- Adicionalmente se reescribe el Excel embebido con ``openpyxl`` para
  garantizar que cualquier apertura manual del archivo muestre los
  mismos datos que el gráfico.

La API pública relevante es :func:`update_chart_with_dataframe`.
"""

from __future__ import annotations

import io
import logging
from typing import Iterable, Optional

import pandas as pd
from openpyxl import Workbook, load_workbook
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn

logger = logging.getLogger(__name__)


class ChartUpdaterError(Exception):
    """Error base del módulo chart_updater."""


class ChartDataShapeError(ChartUpdaterError):
    """El DataFrame no tiene forma compatible para alimentar el gráfico."""


# ---------------------------------------------------------------------- #
# Helpers internos                                                        #
# ---------------------------------------------------------------------- #


def _coerce_numeric(value) -> float:
    """Convierte un valor arbitrario a ``float``. Si no puede, devuelve 0.0."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def _split_dataframe(df: pd.DataFrame) -> tuple[list[str], list[str], list[list[float]]]:
    """Separa un DataFrame en (categorías, nombres de serie, valores).

    - La primera columna del DataFrame se usa como categorías.
    - El resto de las columnas se usan como series.
    - Todos los valores se fuerzan a ``float`` (los no convertibles a 0.0).
    """
    if df is None or df.empty:
        raise ChartDataShapeError(
            "El DataFrame para alimentar el gráfico está vacío."
        )
    if df.shape[1] < 2:
        raise ChartDataShapeError(
            "El DataFrame necesita al menos dos columnas "
            "(una de categorías y una o más de series)."
        )

    categories = [str(v) if v is not None else "" for v in df.iloc[:, 0].tolist()]
    series_names = [str(c) for c in df.columns[1:]]
    series_values: list[list[float]] = []
    for col in df.columns[1:]:
        series_values.append([_coerce_numeric(v) for v in df[col].tolist()])

    return categories, series_names, series_values


def _build_chart_data(df: pd.DataFrame) -> CategoryChartData:
    """Construye un ``CategoryChartData`` de python-pptx a partir del DF."""
    categories, series_names, series_values = _split_dataframe(df)
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in zip(series_names, series_values):
        chart_data.add_series(name, values)
    return chart_data


# ---------------------------------------------------------------------- #
# Actualización del Excel embebido                                        #
# ---------------------------------------------------------------------- #


def _get_embedded_xlsx_part(chart):
    """Devuelve el ``Part`` del Excel embebido asociado al gráfico.

    Recorre las relaciones del part del chart buscando una con tipo
    ``...package`` que corresponde al paquete xlsx embebido. Si no hay
    Excel embebido devuelve ``None``.
    """
    chart_part = chart.part
    package_rel_type = (
        "http://schemas.openxmlformats.org/officeDocument/2006/"
        "relationships/package"
    )
    for rel in chart_part.rels.values():
        if rel.reltype == package_rel_type:
            return rel.target_part
    return None


def _write_dataframe_to_workbook(df: pd.DataFrame) -> bytes:
    """Serializa un DataFrame a un workbook xlsx ``bytes``.

    El workbook tiene una sola hoja ``Sheet1`` con el encabezado en la
    primera fila y los datos en las siguientes. Es la forma que espera
    PowerPoint para los gráficos simples.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Encabezado.
    ws.append([str(c) for c in df.columns.tolist()])
    # Filas. Dejamos la primera columna como texto (categorías) y el
    # resto como números cuando sea posible.
    for row in df.itertuples(index=False, name=None):
        cleaned = [row[0] if row else ""]
        for value in row[1:]:
            cleaned.append(_coerce_numeric(value))
        ws.append(cleaned)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _rewrite_embedded_xlsx(chart, df: pd.DataFrame) -> bool:
    """Reescribe el Excel embebido del gráfico con el contenido del DF.

    Devuelve ``True`` si el Excel embebido fue actualizado, ``False`` si
    el gráfico no tenía Excel embebido (caso raro pero posible).
    """
    embedded_part = _get_embedded_xlsx_part(chart)
    if embedded_part is None:
        logger.debug(
            "El gráfico '%s' no tiene Excel embebido; "
            "se omite la reescritura.",
            getattr(chart, "chart_title", "<sin titulo>"),
        )
        return False

    new_blob = _write_dataframe_to_workbook(df)
    embedded_part._blob = new_blob  # pylint: disable=protected-access
    return True


# ---------------------------------------------------------------------- #
# Búsqueda de gráficos dentro de un slide                                 #
# ---------------------------------------------------------------------- #


def iter_charts_in_slide(slide) -> Iterable[tuple[int, object]]:
    """Itera ``(indice, chart)`` para todos los gráficos de un slide.

    El índice respeta el orden de aparición de las shapes dentro del
    slide, que es el orden en el que ``python-pptx`` las retorna.
    """
    idx = 0
    for shape in slide.shapes:
        if shape.has_chart:
            yield idx, shape.chart
            idx += 1


def find_chart_in_slide(
    slide,
    chart_name: Optional[str] = None,
    chart_index: Optional[int] = None,
):
    """Busca un gráfico en un slide por nombre o por índice.

    Se puede pasar ``chart_name`` (nombre de la shape asociada) o
    ``chart_index`` (posición 0-based entre los gráficos del slide).
    Si se pasan ambos se prioriza ``chart_name`` y ``chart_index`` sirve
    como fallback.

    Devuelve el objeto ``Chart`` de python-pptx o ``None`` si no se
    encuentra.
    """
    charts_with_names: list[tuple[int, object, str]] = []
    idx = 0
    for shape in slide.shapes:
        if shape.has_chart:
            charts_with_names.append((idx, shape.chart, shape.name))
            idx += 1

    if not charts_with_names:
        return None

    if chart_name is not None:
        for _, chart, name in charts_with_names:
            if name == chart_name:
                return chart

    if chart_index is not None:
        if 0 <= chart_index < len(charts_with_names):
            return charts_with_names[chart_index][1]

    if chart_name is None and chart_index is None:
        return charts_with_names[0][1]

    return None


# ---------------------------------------------------------------------- #
# API pública principal                                                   #
# ---------------------------------------------------------------------- #


def update_chart_with_dataframe(chart, df: pd.DataFrame) -> None:
    """Actualiza un gráfico (XML + Excel embebido) con un DataFrame.

    - La primera columna del DF se trata como categorías.
    - El resto de columnas se tratan como series.
    - Se preservan los elementos visuales (colores, fuentes, tamaños)
      porque ``replace_data`` mantiene el estilo del gráfico.
    - Se reescribe el Excel embebido para que el "Edit Data" muestre
      exactamente los mismos valores.

    Parameters
    ----------
    chart:
        Objeto ``pptx.chart.chart.Chart`` obtenido de un slide.
    df:
        DataFrame con los datos. Formato esperado:
        ``[categoria, serie1, serie2, ...]``.

    Raises
    ------
    ChartDataShapeError
        Si el DataFrame no tiene al menos dos columnas o está vacío.
    ChartUpdaterError
        Si ``chart.replace_data`` falla por una shape incompatible.
    """
    chart_data = _build_chart_data(df)

    try:
        chart.replace_data(chart_data)
    except Exception as exc:  # python-pptx lanza varias, las unificamos
        raise ChartUpdaterError(
            f"No se pudo actualizar el gráfico: {exc}"
        ) from exc

    _rewrite_embedded_xlsx(chart, df)


__all__ = [
    "ChartUpdaterError",
    "ChartDataShapeError",
    "update_chart_with_dataframe",
    "find_chart_in_slide",
    "iter_charts_in_slide",
]
